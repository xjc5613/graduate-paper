"""Organize available workbook rows by dates; do not replace original network inputs."""
from pathlib import Path
from datetime import datetime
from collections import Counter
import csv
import hashlib
import json
import re
from openpyxl import load_workbook

SOURCE = Path('专利主题聚类时间all.xlsx')
OUT = Path('data/chapter03')
MAP = {
    '工程软件': {'BIM','参数化设计','生成式设计','数字化设计','智能设计','工程软件','工程仿真模拟'},
    '工程物联网': {'智慧工地','施工监测'},
    '构件智能生产': {'装配式建筑','3D打印'},
    '智能施工机械': {'机器人','盾构机','塔吊','架桥机','智能施工'},
    '工程大数据与互联网平台': {'大数据','大模型','自然语言处理','数字孪生','区块链'},
}

def stage_of(year):
    return 1 if year <= 2014 else 2 if year <= 2019 else 3

def main():
    OUT.mkdir(exist_ok=True)
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb['Sheet1']
    rows, excluded, label = [], [], ''
    for row_num, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row_num in (126, 1711):
            label = ''
        if values[5]:
            label = str(values[5]).strip()
        date = values[3]
        year = date.year if isinstance(date, datetime) else int(str(date)[:4]) if date and re.match(r'^\d{4}', str(date)) else None
        if not values[1] or year is None or not 2010 <= year <= 2025:
            excluded.append(row_num)
            continue
        group = next((key for key, labels in MAP.items() if label in labels), '未提供标签')
        rows.append({'source_row': row_num, 'title': str(values[1]).strip(), 'application_date': date.strftime('%Y-%m-%d') if isinstance(date,datetime) else str(date), 'year': year, 'stage': stage_of(year), 'source_label': label, 'subfield': group, 'ipc': str(values[4] or ''), 'identifier_status': '题名与日期/行号；非申请号'})
    for stage in (1,2,3):
        with (OUT/f'workbook_stage_{stage}_records.csv').open('w', newline='', encoding='utf-8-sig') as f:
            writer=csv.DictWriter(f, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(r for r in rows if r['stage']==stage)
    counts={str(stage):dict(Counter(r['subfield'] for r in rows if r['stage']==stage)) for stage in (1,2,3)}
    unique={str(stage):{key:len({(r['title'],r['year']) for r in rows if r['stage']==stage and r['subfield']==key}) for key in list(MAP)+['未提供标签']} for stage in (1,2,3)}
    with Path('tables/chapter03-annual-from-source-figures.csv').open() as f:
        annual=list(csv.DictReader(f))
    reported={str(stage):{name:sum(int(r[column]) for r in annual if stage_of(int(r['year']))==stage) for name,column in [('patents','patent_count'),('literature','literature_count')]} for stage in (1,2,3)}
    summary={'source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'rows':ws.max_row-1,'eligible_rows':len(rows),'excluded_rows':excluded,'stage_windows':{'1':'2010—2014','2':'2015—2019','3':'2020—2025（2025截至10月1日）'},'reported_annual_aggregates':reported,'workbook_row_counts':counts,'within_subfield_title_year_unique_counts':unique,'mapping':{k:sorted(v) for k,v in MAP.items()},'stage3_complete_year_subfield_unique':{k:len({(r['title'],r['year']) for r in rows if r['subfield']==k and 2020<=r['year']<=2024}) for k in MAP},'stage3_partial_2025_subfield_unique':{k:len({(r['title'],r['year']) for r in rows if r['subfield']==k and r['year']==2025}) for k in MAP},'limits':['前124行无现成标签，未人为填补；缺标签不等于技术未出现','无申请人字段，不能做主体逐阶段统计','CSV为现有工作簿辅助样本集合，不是5926件原网络输入集合','题名年份近似去重、跨类别可重复；不构成专利族或互斥占比','按申请日期重划窗口，不使用工作簿原行区块的阶段年份推断网络窗口']}
    (OUT/'stage-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    assert sum(sum(v.values()) for v in counts.values())==len(rows)==5913
    assert sum(x['patents'] for x in reported.values())==5926
    assert sum(x['literature'] for x in reported.values())==6758
    print(json.dumps({'workbook_stage_rows':{s:sum(v.values()) for s,v in counts.items()},'reported':reported,'source_unchanged':True},ensure_ascii=False))

if __name__=='__main__':
    main()
