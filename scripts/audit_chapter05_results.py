from __future__ import annotations
from pathlib import Path
import ast,csv,json,math,re,hashlib
from collections import defaultdict
from datetime import datetime
import numpy as np
from openpyxl import load_workbook
from zipfile import ZipFile
import xml.etree.ElementTree as E
root=Path('refs/agent-results-2026-10-15')
manifest=json.loads((root/'manifest.json').read_text())
assert all(hashlib.sha256((root/k).read_bytes()).hexdigest()==v for k,v in manifest.items())
tree=ast.parse((root/'code/analyze_forecast.py').read_text())
names={'TOPICS','LABEL_MAP','REGEX_MAP','IPC_MAP','CONVERGENCE','COVERAGE'}
funs={'parse_year','load_patents','patent_key','build_topic_counts','percentile','network_node_scores','weighted_log_slope','project_counts','minmax','score_topics'}
selected=[x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name in funs or isinstance(x,ast.Assign) and any(isinstance(t,ast.Name) and t.id in names for t in x.targets)]
ns={'re':re,'np':np,'math':math,'datetime':datetime,'load_workbook':load_workbook,'INPUTS':root/'original_inputs'}
exec(compile(ast.Module(body=selected,type_ignores=[]),'<audited-calculation-functions>','exec'),ns)
records=ns['load_patents']();counts,topic_records=ns['build_topic_counts'](records)
saved=json.loads((root/'forecast_results.json').read_text());new=ns['score_topics'](counts,saved['network_nodes'])
errors=[]
for row in csv.DictReader((root/'theme_annual_counts.csv').open(encoding='utf-8-sig')):
 for y in range(2015,2026):
  if int(row[str(y)])!=counts[row['topic_id']][y]:errors.append([row['topic_id'],y])
diffs=[]
for t,a in new.items():
 for k in ['balanced_score','market_score','structural_score','log_growth_slope','activity_2021_2026']:
  if abs(a[k]-saved['topics'][t][k])>1e-9:diffs.append([t,k,a[k],saved['topics'][t][k]])
 if {str(k):v for k,v in a['projection'].items()}!=saved['topics'][t]['projection']:diffs.append([t,'projection'])
stages={1:{},2:{},3:{}}
for r in csv.DictReader((root/'ipc_stage_metrics.csv').open(encoding='utf-8-sig')):
 stages[int(r['stage'])][r['ipc']]={k:float(r[k]) for k in ['degree','betweenness','closeness','eigenvector','eccentricity']}
for ipc,v in saved['network_nodes'].items():
 assert all(v[k]==stages[3][ipc][k] for k in ['degree','betweenness','closeness','eigenvector','eccentricity'])
 assert math.isclose(v['network_score'],.6*v['structural']+.4*v['transition'])
common=set(stages[2])&set(stages[3]);delta={k:stages[3][k]['degree']-stages[2][k]['degree']+.7*(stages[3][k]['betweenness']-stages[2][k]['betweenness']) for k in common};vals=sorted(delta.values())
for ipc,v in saved['network_nodes'].items():
 if ipc not in delta: assert v['transition']==.55
 else:
  d=delta[ipc];lo=sum(x<d for x in vals)/(len(vals)-1);hi=(sum(x<=d for x in vals)-1)/(len(vals)-1)
  assert lo-1e-12<=v['transition']<=hi+1e-12
wb=load_workbook(root/'original_inputs/topic_times.xlsx',read_only=True,data_only=True);ws=wb['Sheet1']
print('worksheet rows',ws.max_row,'cols',ws.max_column)
print('boundary rows',[(i, next(ws.iter_rows(min_row=i,max_row=i,values_only=True))[3]) for i in [125,126,1710,1711]])
with ZipFile(root/'original_inputs/ipc_network_results.docx') as z:
 x=E.fromstring(z.read('word/document.xml'));n={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 paras=[''.join(t.text or '' for t in p.findall('.//w:t',n)) for p in x.findall('.//w:p',n)]
 print('stage headings',[s for s in paras if '阶段' in s][:10])
comparisons={}
for a,b in [('original_inputs/manuscript.docx','论文保留初版-基于动态演化的智能建造领域关键技术识别及预测_副本3.docx'),('original_inputs/topic_times.xlsx','专利主题聚类时间all.xlsx'),('original_inputs/ipc_network_results.docx','IPCnetwork结果.docx')]:
 comparisons[a]=Path(b).exists() and hashlib.sha256((root/a).read_bytes()).digest()==hashlib.sha256(Path(b).read_bytes()).digest()
raw_slopes={}
for t in counts:
 y=np.log1p([counts[t][yr] for yr in range(2018,2025)]);x=np.arange(2018,2025)-2021;w=.72**(2024-np.arange(2018,2025));X=np.column_stack([np.ones(7),x]);beta=np.linalg.solve(X.T@(w[:,None]*X),X.T@(w*y));raw_slopes[t]=float(beta[1])
report={'manifest_entries_verified':len(manifest),'source_identity_vs_workspace':comparisons,'labeled_records':len(records),'unique_title_year':len({ns['patent_key'](r) for r in records}),'annual_count_mismatches':errors,'score_projection_mismatches_using_saved_nodes':diffs,'stage_node_counts':{str(k):len(v) for k,v in stages.items()},'raw_slopes_before_clipping':raw_slopes,'counting_rule':'title+year within each theme, cross-theme overlap allowed','convergence':'preassigned constants, not computed from cluster distance matrix','ipc_topic_mapping':'preassigned functional mapping, not document-level membership','rank_tie_implementation':'ordinal ranks; equal values not averaged; transition input assembled from a set; rerun order sensitivity possible','partial_and_estimate_2025_not_added_together':True,'formal_claim_audit_export_in_zip':False}
Path('plan/review/chapter05-results-data-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False,indent=2))
